"""Build a concise Word reference for the current BNPL rule base."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "deliverables" / "BNPL专家系统完整规则简表.docx"


def cell_style(cell, header=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    tc_pr.append(borders)
    margins = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:w"), "85")
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    tc_pr.append(margins)
    if header:
        fill = OxmlElement("w:shd")
        fill.set(qn("w:fill"), "DFEAF4")
        tc_pr.append(fill)
        for run in cell.paragraphs[0].runs:
            run.bold = True


def table(document, headings, rows, widths):
    grid = document.add_table(rows=1, cols=len(headings))
    grid.autofit = False
    for index, width in enumerate(widths):
        grid.columns[index].width = Cm(width)
    header_properties = grid.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for index, heading in enumerate(headings):
        cell = grid.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.text = heading
        cell_style(cell, header=True)
    for row_values in rows:
        row = grid.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.width = Cm(widths[index])
            cell.text = str(value)
            cell_style(cell)
    return grid


def add_heading(document, text):
    document.add_paragraph(text, style="Heading 1")


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    normal = document.styles["Normal"]
    normal.font.name = "Heiti SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 17), ("Heading 1", 11.5)):
        style = document.styles[name]
        style.font.name = "Heiti SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        if name == "Title":
            border = style._element.pPr.find(qn("w:pBdr"))
            if border is not None:
                style._element.pPr.remove(border)

    document.add_paragraph("BNPL专家系统完整规则简表", style="Title")
    document.add_paragraph(
        "这份简表汇总当前可运行版本的先用后付审批与提额规则。系统先执行硬规则，再合成软规则 CF，"
        "按固定梯度得到额度；新用户还须经过信息覆盖率封顶。老用户不适用该封顶。"
    )

    add_heading(document, "输入与适用范围")
    document.add_paragraph(
        "输入字段共 14 个：实名认证、年龄、绑卡类型、注册月数、履约率、月均消费、完成订单数、"
        "连续按时还款次数、近半年失信次数、设备稳定性、当前额度、可用额度、外部信用分、外部逾期记录。"
        "当前额度与可用额度均为空表示新用户；两者均有值表示老用户。新用户不使用履约率、连续还款次数和额度使用率规则。"
    )
    document.add_paragraph(
        "可选字段缺失时不触发相应软规则，贡献为 0；只有已取得且适用的字段进入 CF 分母。"
        "外部逾期记录的“无”表示已确认无逾期，与字段缺失不同。老用户额度使用率 =（当前额度 − 可用额度）÷ 当前额度；"
        "当前额度为 0 时程序按 0 处理。"
    )

    add_heading(document, "硬规则")
    document.add_paragraph("命中任意一条即停止额度计算：新用户拒绝，老用户暂停，最终额度均为 0；不再采用转人工分支。")
    table(document, ("编号", "触发条件", "处理结果"), [
        ("HR-01", "实名认证未通过", "新用户 reject；老用户 suspend"),
        ("HR-02", "年龄低于 18 岁", "新用户 reject；老用户 suspend"),
        ("HR-03", "近半年失信次数 ≥ 2", "新用户 reject；老用户 suspend"),
        ("HR-04", "设备状态为 Abnormal", "新用户 reject；老用户 suspend"),
        ("HR-05", "外部信用分 ≤ 450", "新用户 reject；老用户 suspend"),
    ], (2.3, 7.1, 7.3))

    add_heading(document, "软规则")
    document.add_paragraph("同一字段的分段规则互斥；未列出的区间为中性，不触发规则。表中 CF 是命中时的贡献，不是最终 CF。")
    soft_rules = [
        ("SR-01", "实名通过且绑定信用卡", "+0.45", "全部"),
        ("SR-02", "实名通过且绑定借记卡", "+0.25", "全部"),
        ("SR-03", "未绑定银行卡", "−0.25", "全部"),
        ("SR-04", "注册 ≥ 12 个月", "+0.35", "全部"),
        ("SR-05", "注册 6–11 个月", "+0.20", "全部"),
        ("SR-06", "注册 < 1 个月", "−0.20", "全部"),
        ("SR-07", "履约率 ≥ 98%", "+0.65", "老用户"),
        ("SR-08", "95% ≤ 履约率 < 98%", "+0.40", "老用户"),
        ("SR-09", "90% ≤ 履约率 < 95%", "+0.15", "老用户"),
        ("SR-10", "履约率 < 90%", "−0.50", "老用户"),
        ("SR-11", "月均消费 ≥ 2,000 元", "+0.30", "全部"),
        ("SR-12", "1,000 ≤ 月均消费 < 2,000 元", "+0.15", "全部"),
        ("SR-13", "完成订单数 ≥ 100", "+0.30", "全部"),
        ("SR-14", "完成订单数 50–99", "+0.15", "全部"),
        ("SR-15", "连续按时还款 ≥ 6 次", "+0.55", "老用户"),
        ("SR-16", "连续按时还款 3–5 次", "+0.25", "老用户"),
        ("SR-17", "近半年失信次数 = 1", "−0.35", "全部"),
        ("SR-18", "设备状态为 Stable", "+0.25", "全部"),
        ("SR-19", "设备状态为 Changing", "−0.15", "全部"),
        ("SR-20", "外部信用分 ≥ 800", "+0.55", "全部"),
        ("SR-21", "550 ≤ 外部信用分 < 800", "+0.30", "全部"),
        ("SR-22", "450 < 外部信用分 < 550", "0", "全部"),
        ("SR-23", "外部逾期记录为 Minor", "−0.25", "全部"),
        ("SR-24", "外部逾期记录为 Moderate", "−0.55", "全部"),
        ("SR-25", "外部逾期记录为 Serious", "−0.70", "全部"),
        ("SR-26", "确认无外部逾期记录", "+0.30", "全部"),
        ("SR-27", "40% ≤ 额度使用率 ≤ 80%", "+0.30", "老用户"),
        ("SR-28", "额度使用率 < 20%", "0", "老用户"),
        ("SR-29", "额度使用率 > 90%", "−0.30", "老用户"),
    ]
    table(document, ("编号", "命中条件", "CF", "适用"), soft_rules, (2.2, 9.7, 2.2, 2.6))

    add_heading(document, "CF 合成与额度映射")
    document.add_paragraph(
        "原始 CF = 所有命中软规则的 CF 之和。最终 CF = 原始 CF ÷ 已取得且适用字段的正向 CF 满分，"
        "结果限制在 0–1；分母为 0 时最终 CF 为 0。字段未命中规则时贡献为 0，但只要该字段已取得，"
        "它仍计入正向满分。完整资料下，新用户正向满分为 2.50，老用户为 4.00。"
    )
    table(document, ("最终 CF", "CF 对应额度", "审批结果"), [
        ("< 0.40", "0 元", "不获批"),
        ("0.40 至 < 0.60", "1,000 元", "有条件批准"),
        ("0.60 至 < 0.85", "5,000 元", "批准"),
        ("≥ 0.85", "10,000 元", "批准"),
    ], (5.4, 5.2, 6.1))

    add_heading(document, "新用户信息覆盖率封顶")
    document.add_paragraph(
        "信息覆盖率 = 当前已取得字段的适用正向 CF 满分 ÷ 2.50。该指标只判断正向软规则所需信息是否取得，"
        "不评价字段取值好坏，也不包含仅用于硬规则的字段。封顶额度与 CF 对应额度取较小值；老用户不适用。"
    )
    table(document, ("编号", "信息覆盖率", "最高额度"), [
        ("EC-01", "< 70%", "1,000 元"),
        ("EC-02", "70% 至 < 90%", "5,000 元"),
        ("EC-03", "≥ 90%", "10,000 元"),
    ], (3.0, 6.4, 7.3))

    add_heading(document, "目标额度反向诊断")
    document.add_paragraph(
        "用户可指定 1,000、5,000 或 10,000 元目标。系统先检查硬规则和现有结果；新用户若被信息覆盖率封顶，"
        "先说明封顶并列出缺失的软规则字段。否则计算达到目标 CF 的差距，并分别模拟注册月数达到 6 或 12、"
        "订单数达到 50 或 100、老用户连续按时还款达到 3 或 6 次，以及通过偿还欠款把额度使用率降至 80% 或以下。"
        "每次只改变一个字段，并重新执行原规则；没有单项方案达标时直接说明。模拟不是提额保证，也不是 CLIPS 原生反向链。"
    )
    document.add_paragraph(
        "说明：70% 与 90% 是课程项目政策阈值，尚未用真实违约数据校准；本简表与当前源代码同步，"
        "适用于课堂原型，不应直接用于真实授信。"
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
