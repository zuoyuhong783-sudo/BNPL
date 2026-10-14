"""Create the project rule amendment as a Word document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "新用户信息覆盖率额度封顶规则.docx"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        tag = "w:" + side
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")


def set_cell_margins(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "start", "bottom", "end"):
        element = OxmlElement("w:" + side)
        element.set(qn("w:w"), "90")
        element.set(qn("w:type"), "dxa")
        margins.append(element)


def add_table(document, rows):
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Cm(5.0), Cm(4.0), Cm(6.6)]
    for index, label in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        cell.text = label
        set_cell_fill(cell, "DCE9F5")
        set_cell_border(cell)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row_values in rows[1:]:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.width = widths[index]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
    return table


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = "Heiti SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 18), ("Heading 1", 12)):
        style = document.styles[name]
        style.font.name = "Heiti SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        if name == "Title":
            paragraph_properties = style._element.pPr
            border = paragraph_properties.find(qn("w:pBdr"))
            if border is not None:
                paragraph_properties.remove(border)

    document.add_paragraph("新用户信息覆盖率额度封顶规则", style="Title")
    document.add_paragraph(
        "本文件补充现有先用后付专家系统的额度映射规则。新用户的 CF 计算方式和原有硬规则、"
        "软规则均不变；系统在算出 CF 对应额度后，再按已取得的正向软规则信息覆盖率确定最高档位。"
        "最终额度取两者较低值。该政策只适用于新用户申请，不用于老用户提额审核。"
    )

    document.add_paragraph("规则定义", style="Heading 1")
    document.add_paragraph(
        "信息覆盖率 = 当前已取得字段对应的适用正向 CF 满分 ÷ 新用户全部正向 CF 满分 2.50。"
        "它衡量信息是否可用，不衡量字段取值是否良好。例如，确认有外部逾期记录与确认无逾期记录，"
        "都表示该字段已取得；完全未取得才不计入覆盖率。缺失字段仍按原规则不贡献 CF。"
    )
    add_table(document, [
        ("CLIPS 规则", "信息覆盖率", "新用户最高额度"),
        ("EC-01", "低于 70%", "1,000 元"),
        ("EC-02", "70% 至低于 90%", "5,000 元"),
        ("EC-03", "90% 及以上", "10,000 元，不额外压低"),
    ])

    document.add_paragraph("决策顺序", style="Heading 1")
    document.add_paragraph(
        "第一步执行硬规则。任一硬规则触发时，新用户拒绝并将额度设为 0，不计算 CF 或覆盖率封顶。"
        "第二步按原有公式计算 CF，即触发的软规则 CF 之和除以适用正向 CF 满分，结果小于 0 按 0 计。"
        "第三步用既有梯度得到 CF 对应额度 0、1,000、5,000 或 10,000 元。"
        "第四步执行 EC-01 至 EC-03 中唯一一条封顶规则，最终额度取 CF 对应额度和封顶额度的较小值。"
        "封顶只能降低额度，不能把原本不获批的用户提高到 1,000 元。"
    )

    document.add_paragraph("算例与反向诊断", style="Heading 1")
    document.add_paragraph(
        "一位新用户已实名、绑定信用卡、注册满 12 个月且设备稳定，但没有提供消费、订单与外部信用数据。"
        "现有软规则得分为 1.05，适用正向满分也是 1.05，因此 CF = 1.00，原 CF 档位为 10,000 元。"
        "信息覆盖率为 1.05 ÷ 2.50 = 42%，命中 EC-01，最终额度为 1,000 元。"
    )
    document.add_paragraph(
        "当目标额度高于当前封顶时，反向诊断先说明封顶原因并列出缺失的软规则输入字段。"
        "补充资料后必须重新执行全部规则；补齐字段不保证达到目标，且不能虚构数据或承诺立即提额。"
    )

    document.add_paragraph("适用边界", style="Heading 1")
    document.add_paragraph(
        "70% 与 90% 是课程项目设定的政策阈值，尚未经真实违约数据校准。覆盖率只包含正向软规则字段，"
        "不涵盖仅用于硬规则的字段，例如近半年失信次数；这些字段若缺失，应在未来的数据质量方案中另行处理。"
        "老用户本版不加覆盖率封顶，避免外部数据暂时缺失时自动压低现有额度。"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
